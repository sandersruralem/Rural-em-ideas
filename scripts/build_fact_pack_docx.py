#!/usr/bin/env python3
"""Render the office-oriented fact-pack Markdown as a styled Word document."""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is required: python3 -m pip install -r requirements-docx.txt")


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "context" / "sanders-county-fact-pack.docx-source.md"
OUTPUT = ROOT / "artifacts" / "Sanders-County-Fact-Pack-DRAFT.docx"

NAVY = "17365D"
BLUE = "2F5597"
PALE_BLUE = "D9EAF7"
PALE_GOLD = "FFF2CC"
LIGHT_GRAY = "E7E6E6"
DARK_GRAY = "404040"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    """Add simple bold/italic Markdown and active URLs."""
    text = text.replace(r"\|", "|")
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+\*|https?://[^\s<>]+|<br>)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token == "<br>":
            paragraph.add_run().add_break()
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("http"):
            url = token.rstrip(".,;)")
            add_hyperlink(paragraph, url, url)
            if len(token) > len(url):
                paragraph.add_run(token[len(url) :])
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        ("Title", 26, NAVY),
        ("Heading 1", 17, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 10.5, BLUE),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    if "Subtitle" not in styles:
        styles.add_style("Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos Display"
    subtitle.font.size = Pt(18)
    subtitle.font.color.rgb = RGBColor.from_string(BLUE)
    subtitle.paragraph_format.space_after = Pt(18)

    for list_style in ("List Bullet", "List Number"):
        styles[list_style].font.name = "Aptos"
        styles[list_style].font.size = Pt(9.5)
        styles[list_style].paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.text = "SANDERS COUNTY, MONTANA  |  COUNTY FACT PACK"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(BLUE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("DRAFT • GREEN DATA • VERIFY BEFORE OPERATIONAL USE   |   ")
    run.font.name = "Aptos"
    run.font.size = Pt(7.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    add_page_number(footer)

    core = document.core_properties
    core.title = "County Fact Pack — Sanders County, Montana"
    core.subject = "Emergency management local-context research draft"
    core.author = "Sanders County Emergency Management"
    core.keywords = "Sanders County, emergency management, fact pack, public information"


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].startswith("|"):
        row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        del rows[1]
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            value = values[column_index] if column_index < len(values) else ""
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8 if column_count >= 4 else 8.5)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif "UNKNOWN" in value.upper() or "VERIFY" in value.upper():
                set_cell_shading(cell, PALE_GOLD)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F5F7FA")
        if row_index == 0:
            set_repeat_table_header(row)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_notice(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GOLD)
    set_cell_margins(cell, 130, 150, 130, 150)
    paragraph = cell.paragraphs[0]
    add_inline(paragraph, text.lstrip("> ").strip())
    paragraph.paragraph_format.space_after = Pt(0)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body(document: Document, lines: list[str]) -> None:
    index = 0
    first_h1 = True
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line[2:])
            first_h1 = False
        elif line == "## Sanders County, Montana" and not first_h1:
            paragraph = document.add_paragraph(style="Subtitle")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line[3:])
        elif line.startswith("## "):
            title = line[3:]
            if title.startswith("1.") or title.startswith("Open items") or title.startswith("Appendix"):
                document.add_page_break()
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, title)
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:])
        elif line.startswith("> "):
            quote_lines = [line[2:]]
            index += 1
            while index < len(lines) and lines[index].startswith(">"):
                quote_lines.append(lines[index].lstrip("> ").strip())
                index += 1
            add_notice(document, " ".join(quote_lines))
            continue
        elif line.startswith("- "):
            parts = [line[2:].strip()]
            index += 1
            while index < len(lines) and lines[index].startswith("  "):
                parts.append(lines[index].strip())
                index += 1
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, " ".join(parts))
            continue
        elif re.match(r"\d+\.\s+", line):
            parts = [re.sub(r"^\d+\.\s+", "", line)]
            index += 1
            while index < len(lines) and lines[index].startswith("   "):
                parts.append(lines[index].strip())
                index += 1
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, " ".join(parts))
            continue
        else:
            parts = [line.strip()]
            index += 1
            while (
                index < len(lines)
                and lines[index].strip()
                and not re.match(r"^(#|\||- |\d+\. |> )", lines[index])
            ):
                parts.append(lines[index].strip())
                index += 1
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(parts))
            continue
        index += 1


def main() -> None:
    if not SOURCE.exists():
        sys.exit(f"Missing source file: {SOURCE}")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_body(document, SOURCE.read_text(encoding="utf-8").splitlines())
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
